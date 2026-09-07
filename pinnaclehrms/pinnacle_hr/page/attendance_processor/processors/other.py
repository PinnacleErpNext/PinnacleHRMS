import io
from datetime import datetime, time

import frappe
from openpyxl import load_workbook


def process(file):
    """
    Process generic attendance Excel file.

    Expected columns:

        Employee
        Employee Name
        Attendance Date
        In Time
        Out Time
    """

    file_stream = file.stream.read()

    wb = load_workbook(
        filename=io.BytesIO(file_stream),
        data_only=True,
    )

    sheet = wb.active

    if sheet.max_row < 2:
        return []

    headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]

    col_index = {header: index for index, header in enumerate(headers)}

    required = [
        "Employee",
        "Employee Name",
        "Attendance Date",
        "In Time",
        "Out Time",
    ]

    for field in required:
        if field not in col_index:
            frappe.throw(f"Missing required column: {field}")

    def format_date(value):
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")

        for fmt in (
            "%Y-%m-%d",
            "%d/%m/%Y",
            "%m/%d/%Y",
            "%d-%b-%Y",
        ):
            try:
                return datetime.strptime(
                    str(value),
                    fmt,
                ).strftime("%Y-%m-%d")
            except ValueError:
                continue

        return str(value)

    def format_time(value):
        if value in (
            None,
            "",
            " ",
        ):
            return ""

        if isinstance(
            value,
            (datetime, time),
        ):
            return value.strftime("%H:%M:%S")

        return str(value).strip()

    records = []
    seen = set()

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        try:
            employee = row[col_index["Employee"]]

            employee_name = row[col_index["Employee Name"]]

            date_value = row[col_index["Attendance Date"]]

            in_value = row[col_index["In Time"]]

            out_value = row[col_index["Out Time"]]

            if not employee or not employee_name or not date_value:
                continue

            attendance_date = format_date(date_value)

            in_time = format_time(in_value)

            out_time = format_time(out_value)

            key = (
                str(employee),
                attendance_date,
                in_time,
                out_time,
            )

            if key in seen:
                continue

            seen.add(key)

            if in_time:
                records.append(
                    {
                        "employee": str(employee),
                        "employee_id": str(employee),
                        "employee_name": str(employee_name).strip(),
                        "device_id": "",
                        "device": "Other",
                        "device_name": "Other",
                        "attendance_date": attendance_date,
                        "time": in_time,
                        "source": "Other",
                    }
                )

            if out_time:
                records.append(
                    {
                        "employee": str(employee),
                        "employee_id": str(employee),
                        "employee_name": str(employee_name).strip(),
                        "device_id": "",
                        "device": "Other",
                        "device_name": "Other",
                        "attendance_date": attendance_date,
                        "time": out_time,
                        "source": "Other",
                    }
                )

        except Exception as exc:
            frappe.log_error(
                f"Error processing Other attendance row: {exc}",
                "Attendance Formatter",
            )

    return records
