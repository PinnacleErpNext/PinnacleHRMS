import io
from datetime import datetime, time

import frappe
from openpyxl import load_workbook

from ..utils.normalization import merge_header_cells


def process(file):
    """
    Process Mantra attendance Excel file.
    """

    file_stream = file.stream.read()

    wb = load_workbook(
        filename=io.BytesIO(file_stream),
        data_only=True,
    )

    sheet = wb.active

    raw_header = [cell.value for cell in sheet[1]]

    header = merge_header_cells(raw_header)

    header = [str(value).lower().strip() for value in header]

    col_idx = {value: index for index, value in enumerate(header)}

    required = [
        "attendance device id",
        "attendance device",
        "employee name",
        "attendance date",
        "in time",
        "out time",
    ]

    for field in required:
        if field not in col_idx:
            frappe.throw(f"Missing required column: {field}")

    records = []
    seen = set()

    def format_date(value):
        if isinstance(value, datetime):
            return value.strftime("%Y-%m-%d")

        for fmt in (
            "%m/%d/%Y",
            "%d/%m/%Y",
            "%Y-%m-%d",
            "%d-%b-%Y",
        ):
            try:
                return datetime.strptime(
                    str(value),
                    fmt,
                ).strftime("%Y-%m-%d")
            except ValueError:
                continue

        return str(value)

    def format_time(value):
        if value in (
            None,
            "",
            " ",
        ):
            return ""

        if isinstance(
            value,
            (datetime, time),
        ):
            return value.strftime("%H:%M:%S")

        return str(value).strip()

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        device_id = str(row[col_idx["attendance device id"]] or "").strip()

        device = str(row[col_idx["attendance device"]] or "").strip()

        employee_name = str(row[col_idx["employee name"]] or "").strip()

        date_value = row[col_idx["attendance date"]]

        in_value = row[col_idx["in time"]]

        out_value = row[col_idx["out time"]]

        if not device_id or not employee_name or not date_value:
            continue

        attendance_date = format_date(date_value)

        in_time = format_time(in_value)

        out_time = format_time(out_value)

        key = (
            device_id,
            employee_name,
            attendance_date,
            in_time,
            out_time,
        )

        if key in seen:
            continue

        seen.add(key)

        # Add IN as a raw punch
        if in_time:
            records.append(
                {
                    "employee": None,
                    "employee_id": None,
                    "device_id": device_id,
                    "device": device,
                    "device_name": device,
                    "employee_name": employee_name,
                    "attendance_date": attendance_date,
                    "time": in_time,
                    "source": "Mantra",
                }
            )

        # Add OUT as a raw punch
        if out_time:
            records.append(
                {
                    "employee": None,
                    "employee_id": None,
                    "device_id": device_id,
                    "device": device,
                    "device_name": device,
                    "employee_name": employee_name,
                    "attendance_date": attendance_date,
                    "time": out_time,
                    "source": "Mantra",
                }
            )

    frappe.msgprint(f"✅ Processed {len(records)} Mantra records")

    return records
