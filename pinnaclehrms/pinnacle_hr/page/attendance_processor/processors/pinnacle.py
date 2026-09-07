import io
import re
from datetime import datetime

import frappe
from openpyxl import load_workbook

from ..constants import PINNACLE_DEVICE_NAME

TIME_PATTERN = re.compile(r"\d{2}:\d{2}")


def process(file):
    """
    Process Pinnacle biometric Excel file.

    Expected sheet:
        Att.log report
    """

    file_stream = file.stream.read()

    wb = load_workbook(
        filename=io.BytesIO(file_stream),
        data_only=True,
    )

    if "Att.log report" not in wb.sheetnames:
        frappe.throw("Sheet 'Att.log report' not found.")

    ws = wb["Att.log report"]

    raw_period = ws["C3"].value

    if not raw_period:
        frappe.throw("Payroll period not found in Pinnacle file.")

    start_date_str = str(raw_period).split("~")[0].strip()

    try:
        start_date = datetime.strptime(
            start_date_str,
            "%Y-%m-%d",
        )
    except ValueError:
        frappe.throw(f"Invalid Pinnacle period format: {raw_period}")

    formatted_period = start_date.strftime("%b-%Y")

    # Day numbers from row 4.
    # Preserve the actual Excel column index.
    date_columns = []

    for cell in ws[4]:
        if isinstance(cell.value, int):
            date_columns.append((cell.column, cell.value))

    records = []

    row = 5
    max_row = ws.max_row

    while row <= max_row:

        if (
            ws.cell(
                row=row,
                column=1,
            ).value
            == "ID:"
        ):

            device_id = str(
                ws.cell(
                    row=row,
                    column=3,
                ).value
                or ""
            ).strip()

            employee_name = str(
                ws.cell(
                    row=row,
                    column=11,
                ).value
                or ""
            ).strip()

            punch_row = row + 1

            for col_index, day in date_columns:

                punch_cell = ws.cell(
                    row=punch_row,
                    column=col_index,
                ).value

                if not punch_cell:
                    continue

                if not isinstance(
                    punch_cell,
                    str,
                ):
                    punch_cell = str(punch_cell)

                times = TIME_PATTERN.findall(punch_cell.replace(" ", ""))

                if not times:
                    continue

                try:
                    attendance_date = datetime.strptime(
                        f"{day:02d}-{formatted_period}",
                        "%d-%b-%Y",
                    ).date()
                except Exception:
                    continue

                for punch_time in times:

                    records.append(
                        {
                            "employee": None,
                            "employee_id": None,
                            "device_id": device_id,
                            "employee_name": employee_name,
                            "device": PINNACLE_DEVICE_NAME,
                            "device_name": PINNACLE_DEVICE_NAME,
                            "attendance_date": attendance_date,
                            "time": punch_time,
                            "source": "Pinnacle",
                        }
                    )

            row += 2

        else:
            row += 1

    return records
