import io
from datetime import datetime

import frappe
from openpyxl import load_workbook

from ..constants import OPTICODE_DEVICE_NAME
from ..utils.date_time import parse_time_safe


def process(file):
    """
    Process Opticode Final Excel file.

    Expected sheet:
        Final
    """

    file_stream = file.stream.read()

    wb = load_workbook(
        filename=io.BytesIO(file_stream),
        data_only=True,
    )

    if "Final" not in wb.sheetnames:
        frappe.throw("Sheet 'Final' not found in the workbook.")

    sheet = wb["Final"]

    if sheet.max_row < 2:
        return []

    headers = [str(cell.value).strip() if cell.value else "" for cell in sheet[1]]

    col = {header: index for index, header in enumerate(headers)}

    required = [
        "ID",
        "G",
        "Date",
        "In Time",
        "Out Time",
    ]

    for field in required:
        if field not in col:
            frappe.throw(f"Missing required column: {field}")

    records = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):

        try:
            device_id = row[col["ID"]]
            employee_name = row[col["G"]]
            raw_date = row[col["Date"]]

            in_value = row[col["In Time"]]
            out_value = row[col["Out Time"]]

            if not device_id or not employee_name or not raw_date:
                continue

            # Date
            try:
                attendance_date = (
                    raw_date.date()
                    if isinstance(
                        raw_date,
                        datetime,
                    )
                    else raw_date
                )
            except Exception:
                continue

            # IN
            in_time = parse_time_safe(in_value)

            if in_time:
                records.append(
                    {
                        "employee": None,
                        "employee_id": None,
                        "device_id": str(device_id),
                        "employee_name": str(employee_name).strip(),
                        "device": OPTICODE_DEVICE_NAME,
                        "device_name": OPTICODE_DEVICE_NAME,
                        "attendance_date": attendance_date,
                        "time": in_time,
                        "source": "Opticode",
                    }
                )

            # OUT
            out_time = parse_time_safe(out_value)

            if out_time:
                records.append(
                    {
                        "employee": None,
                        "employee_id": None,
                        "device_id": str(device_id),
                        "employee_name": str(employee_name).strip(),
                        "device": OPTICODE_DEVICE_NAME,
                        "device_name": OPTICODE_DEVICE_NAME,
                        "attendance_date": attendance_date,
                        "time": out_time,
                        "source": "Opticode",
                    }
                )

        except Exception as exc:
            frappe.log_error(
                f"Opticode Final processing error: {exc}",
                "Opticode Formatter",
            )

    return records
